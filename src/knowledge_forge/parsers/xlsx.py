import asyncio
import io
import logging

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from knowledge_forge.collectors.base import RawItem
from knowledge_forge.exceptions import ParsingError
from knowledge_forge.ids import stable_id
from knowledge_forge.models.document import Document, DocumentElement, ImageElement, TableElement, TextBlock
from knowledge_forge.models.source import SourceRef
from knowledge_forge.parsers.base import Parser
from knowledge_forge.vision.extractor import VisionExtractor
from knowledge_forge.vision.factory import get_vision_extractor

logger = logging.getLogger(__name__)

# A block with fewer rows/columns than this doesn't look like a data grid -
# it's classified as a title (first block on the sheet) or a note (any
# other block), never a table. See _classify_block.
_MIN_TABLE_HEIGHT = 2
_MIN_TABLE_WIDTH = 2
_NOTE_PREFIXES = ("※", "*", "주:", "주 :", "비고", "Note", "note:", "NOTE")
_MAX_HEADER_ROWS = 3


class XlsxParser(Parser):
    """Parses .xlsx into a Document using openpyxl.

    Each worksheet is split into vertically-separated blocks (runs of
    non-blank rows) and each block becomes either a TableElement (a real
    data grid) or a TextBlock (a title above the grid, or a footnote below
    it) - see _split_blocks/_classify_block. This keeps a sheet's title and
    footnotes out of the table's own rows, which would otherwise corrupt
    row-aware chunking (chunking.hierarchy.split_table_by_rows).

    Within a table block, merged cells are propagated (the top-left value
    fills every cell in the merge, matching how python-docx/pptx already
    surface merged table cells - see docx.py/pptx.py) and consecutive
    bold-styled leading rows are detected as `header_rows`, letting a
    multi-level XLSX header (e.g. a merged "2025년" row above "기본급"/
    "가산율") survive instead of being flattened into one header row. When
    header_rows > 1, a nested `structured` view of the table is also built
    (see _build_structured) for future graph/ontology consumers - chunking
    itself never reads it, only `rows`.

    Embedded images anchored over/in a cell go through VisionExtractor,
    same policy and per-image degradation as DOCX/PPTX/HWPX.

    XLSX has no page concept - every element's `page` stays None.
    """

    supported_extensions = {".xlsx"}

    def __init__(self, vision_extractor: VisionExtractor | None = None):
        self._vision = vision_extractor

    def _vision_extractor(self) -> VisionExtractor:
        if self._vision is None:
            self._vision = get_vision_extractor()
        return self._vision

    async def parse(self, item: RawItem, source_ref: SourceRef) -> Document:
        try:
            elements: list[DocumentElement] = []
            workbook = openpyxl.load_workbook(io.BytesIO(item.content), data_only=True)
            # A second load (data_only=False) is the only way openpyxl
            # exposes a cell's original formula string - `workbook` above
            # gives the cached computed value instead, which is what
            # `rows` needs for rendering/search.
            formula_workbook = openpyxl.load_workbook(io.BytesIO(item.content), data_only=False)
            for sheet in workbook.worksheets:
                formula_sheet = formula_workbook[sheet.title]
                elements.extend(self._parse_sheet(sheet, formula_sheet, item.uri, len(elements)))
                elements.extend(await self._convert_images(sheet, item.uri, len(elements)))
        except Exception as exc:  # noqa: BLE001 - surfaced as a typed parsing error
            raise ParsingError(f"Failed to parse xlsx '{item.uri}': {exc}") from exc

        return Document(
            document_id=stable_id("document", item.uri, item.content_hash),
            source=source_ref,
            elements=elements,
        )

    def _parse_sheet(
        self, sheet: Worksheet, formula_sheet: Worksheet, uri: str, start_index: int
    ) -> list[DocumentElement]:
        values, bold = _build_grid(sheet)
        formula_values = _build_formula_grid(formula_sheet)

        elements: list[DocumentElement] = []
        caption_used = False
        for block_index, row_idxs in enumerate(_split_blocks(values)):
            if _classify_block(row_idxs, values, is_first_block=(block_index == 0)) == "table":
                header_rows = _detect_header_rows(row_idxs, bold)
                caption = sheet.title if not caption_used else None
                caption_used = True
                elements.append(
                    _build_table_element(
                        row_idxs, values, formula_values, header_rows, caption, uri, start_index + len(elements)
                    )
                )
            else:
                elements.append(_build_text_block(row_idxs, values, uri, start_index + len(elements)))
        return elements

    async def _convert_images(self, sheet, uri: str, start_index: int) -> list[DocumentElement]:
        images = list(sheet._images)  # noqa: SLF001 - openpyxl exposes no public accessor

        async def _describe(image) -> str | None:
            try:
                image_bytes = image._data()  # noqa: SLF001 - openpyxl exposes no public accessor
                return await self._vision_extractor().describe_figure(
                    image_bytes, media_type=f"image/{image.format}"
                ) or None
            except Exception as exc:  # noqa: BLE001 - degrade this image, not the whole document
                logger.warning("VLM figure description failed on %s sheet %s: %s", uri, sheet.title, exc)
                return None

        descriptions = await asyncio.gather(*(_describe(image) for image in images))

        elements: list[DocumentElement] = []
        for description in descriptions:
            elements.append(
                ImageElement(
                    element_id=stable_id("element", uri, str(start_index + len(elements)), "image"),
                    ocr_text=description,
                    extraction_method="vlm" if description is not None else None,
                )
            )
        return elements


def _build_grid(sheet: Worksheet) -> tuple[list[list[str]], list[list[bool]]]:
    """Reads the sheet into a value grid and a parallel bold-flag grid, with
    merged ranges propagated - every cell in a merge gets the top-left
    cell's value and bold flag, instead of the None/False openpyxl leaves in
    a merge's non-top-left cells. Without this, a merged header like
    "2025년" spanning B1:C1 would leave C1 blank and silently drop the
    column's label.
    """
    max_row, max_col = sheet.max_row, sheet.max_column
    values = [
        ["" if (v := sheet.cell(r, c).value) is None else str(v) for c in range(1, max_col + 1)]
        for r in range(1, max_row + 1)
    ]
    bold = [
        [bool(sheet.cell(r, c).font.bold) for c in range(1, max_col + 1)] for r in range(1, max_row + 1)
    ]
    for merged_range in sheet.merged_cells.ranges:
        top_value = values[merged_range.min_row - 1][merged_range.min_col - 1]
        top_bold = bold[merged_range.min_row - 1][merged_range.min_col - 1]
        for r in range(merged_range.min_row, merged_range.max_row + 1):
            for c in range(merged_range.min_col, merged_range.max_col + 1):
                values[r - 1][c - 1] = top_value
                bold[r - 1][c - 1] = top_bold
    return values, bold


def _build_formula_grid(sheet: Worksheet) -> list[list[object]]:
    """Raw cell values from a data_only=False load: a formula cell's value
    is its formula string (e.g. "=SUM(B5:B6)"); a merge's non-top-left
    cells and cells with no formula come through as None/a literal.
    """
    return [
        [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)] for r in range(1, sheet.max_row + 1)
    ]


def _split_blocks(values: list[list[str]]) -> list[list[int]]:
    """Groups row indices into contiguous runs separated by fully-blank
    rows - each run is one region (title/table/note) to classify and parse
    independently, instead of treating the whole sheet as one grid.
    """
    blocks: list[list[int]] = []
    current: list[int] = []
    for i, row in enumerate(values):
        if any(cell.strip() for cell in row):
            current.append(i)
        elif current:
            blocks.append(current)
            current = []
    if current:
        blocks.append(current)
    return blocks


def _classify_block(row_idxs: list[int], values: list[list[str]], is_first_block: bool) -> str:
    """Classifies a block as "table" (a grid-shaped block: >=2 rows, >=2
    used columns, not footnote-prefixed), else "title" for the sheet's
    first block, else "note". A single-row or single-column block is never
    a data table - a real table always has a header plus at least one data
    row/column.
    """
    block_rows = [values[i] for i in row_idxs]
    num_cols = len(block_rows[0]) if block_rows else 0
    used_cols = [c for c in range(num_cols) if any(row[c].strip() for row in block_rows)]
    # _split_blocks only ever includes rows with some non-blank cell, so the
    # block's first row always has content.
    first_text = _row_text(block_rows[0]) if block_rows else ""
    is_table_shaped = len(row_idxs) >= _MIN_TABLE_HEIGHT and len(used_cols) >= _MIN_TABLE_WIDTH
    if is_table_shaped and not first_text.startswith(_NOTE_PREFIXES):
        return "table"
    return "title" if is_first_block else "note"


def _detect_header_rows(row_idxs: list[int], bold: list[list[bool]]) -> int:
    """Counts consecutive leading rows (within the block) that carry any
    bold cell, stopping as soon as a row has none - a bold row typically
    marks a header level (openpyxl's default cell style is never bold), and
    stopping guarantees at least one data row remains below the header.
    Falls back to 1 (a single header row) when no bold signal is present at
    all, matching the assumption every other parser already makes.
    """
    count = 0
    for i in row_idxs[:-1]:
        if not any(bold[i]):
            break
        count += 1
        if count >= _MAX_HEADER_ROWS:
            break
    return count or 1


def _row_text(row: list[str]) -> str:
    # Merge propagation (_build_grid) fills every cell in a merged range
    # with the same value - collapse consecutive repeats so a merged title/
    # note spanning N columns renders once, not N times.
    parts: list[str] = []
    for cell in row:
        if cell.strip() and (not parts or parts[-1] != cell):
            parts.append(cell)
    return " ".join(parts)


def _build_text_block(row_idxs: list[int], values: list[list[str]], uri: str, index: int) -> TextBlock:
    lines = [_row_text(values[i]) for i in row_idxs]
    return TextBlock(
        element_id=stable_id("element", uri, str(index), "text"),
        text="\n".join(line for line in lines if line),
    )


def _build_table_element(
    row_idxs: list[int],
    values: list[list[str]],
    formula_values: list[list[object]],
    header_rows: int,
    caption: str | None,
    uri: str,
    index: int,
) -> TableElement:
    rows = [values[i] for i in row_idxs]
    structured = _build_structured(rows[:header_rows], rows[header_rows:]) if header_rows > 1 else None
    formulas = _extract_formulas(row_idxs, formula_values)
    return TableElement(
        element_id=stable_id("element", uri, str(index), "table"),
        rows=rows,
        caption=caption,
        extraction_method="rule_based",
        header_rows=header_rows,
        structured=structured,
        formulas=formulas or None,
    )


def _build_structured(header_block: list[list[str]], data_rows: list[list[str]]) -> dict:
    """Turns a multi-level header + data rows into a nested-dict view: each
    data row becomes a record keyed by its leading column (the row label,
    e.g. "외래"), with the remaining columns nested per header level (e.g.
    record["2025년"]["기본급"]) instead of one flat "2025년/기본급" column.
    """
    row_key = header_block[-1][0].strip() or "항목"
    records = []
    for row in data_rows:
        record: dict = {row_key: row[0]}
        for c in range(1, len(row)):
            path = [h[c].strip() for h in header_block if h[c].strip()]
            if not path:
                continue
            node = record
            for key in path[:-1]:
                node = node.setdefault(key, {})
            node[path[-1]] = row[c]
        records.append(record)
    return {"headers": header_block, "rows": records}


def _extract_formulas(row_idxs: list[int], formula_values: list[list[object]]) -> dict[str, str]:
    formulas = {}
    for i in row_idxs:
        for c, value in enumerate(formula_values[i]):
            if isinstance(value, str) and value.startswith("="):
                formulas[f"{get_column_letter(c + 1)}{i + 1}"] = value
    return formulas
