import asyncio
import io
import logging

from docx import Document as DocxDocument
from docx.document import Document as DocxDocumentType
from docx.enum.shape import WD_INLINE_SHAPE_TYPE

from knowledge_forge.collectors.base import RawItem
from knowledge_forge.exceptions import ParsingError
from knowledge_forge.ids import stable_id
from knowledge_forge.models.document import (
    Document,
    DocumentElement,
    Heading,
    ImageElement,
    ListElement,
    TableElement,
    TextBlock,
)
from knowledge_forge.models.source import SourceRef
from knowledge_forge.parsers.base import Parser
from knowledge_forge.vision.factory import VisionExtractorLike, get_vision_extractor

logger = logging.getLogger(__name__)

# Word's "Increase Indent" button steps left_indent by 0.5in (457200 EMU) -
# used to bucket a paragraph's direct left-indent into a discrete nesting
# level rather than reasoning in raw EMU.
_INDENT_STEP_EMU = 457200


def _is_list_style(style_name: str) -> bool:
    # Matches nested list styles too ("List Bullet 2", "List Number 3", ...),
    # not just the top-level "List Bullet"/"List Number".
    return style_name.startswith("List Bullet") or style_name.startswith("List Number")


def _indent_level(paragraph) -> int | None:
    left_indent = paragraph.paragraph_format.left_indent
    if not left_indent:
        return None
    level = round(left_indent / _INDENT_STEP_EMU)
    return level if level >= 1 else None


class DocxParser(Parser):
    """Parses .docx into a Document using python-docx.

    DOCX has no page concept at the XML level (pagination is a rendering
    detail), so every element's `page` stays None. Tables give exact cell
    text deterministically (no VLM needed). Embedded pictures go through
    VisionExtractor, same policy as PPTX's Picture shapes.
    """

    supported_extensions = {".docx"}

    def __init__(self, vision_extractor: VisionExtractorLike | None = None):
        self._vision = vision_extractor

    def _vision_extractor(self) -> VisionExtractorLike:
        if self._vision is None:
            self._vision = get_vision_extractor()
        return self._vision

    async def parse(self, item: RawItem, source_ref: SourceRef) -> Document:
        try:
            docx_document = DocxDocument(io.BytesIO(item.content))
            elements = self._convert_paragraphs(docx_document, item.uri)
            elements.extend(self._convert_tables(docx_document, item.uri, len(elements)))
            elements.extend(await self._convert_images(docx_document, item.uri, len(elements)))
        except Exception as exc:  # noqa: BLE001 - surfaced as a typed parsing error
            raise ParsingError(f"Failed to parse docx '{item.uri}': {exc}") from exc

        return Document(
            document_id=stable_id("document", item.uri, item.content_hash),
            source=source_ref,
            elements=elements,
        )

    def _convert_paragraphs(self, docx_document: DocxDocumentType, uri: str) -> list[DocumentElement]:
        elements: list[DocumentElement] = []
        pending_list_items: list[str] = []
        pending_list_ordered = False
        pending_list_indent: int | None = None

        def flush_list() -> None:
            nonlocal pending_list_items, pending_list_ordered, pending_list_indent
            if pending_list_items:
                elements.append(
                    ListElement(
                        element_id=stable_id("element", uri, str(len(elements)), "list"),
                        items=pending_list_items,
                        ordered=pending_list_ordered,
                        indent_level=pending_list_indent,
                    )
                )
                pending_list_items = []
                pending_list_indent = None

        for paragraph in docx_document.paragraphs:
            style_name = paragraph.style.name if paragraph.style else "Normal"
            text = paragraph.text.strip()

            if _is_list_style(style_name):
                if not text:
                    continue
                ordered = style_name.startswith("List Number")
                indent = _indent_level(paragraph)
                # A change in list type or nesting depth starts a new
                # ListElement so each one carries a single indent_level,
                # rather than collapsing sub-items into their parent's text.
                if pending_list_items and (ordered != pending_list_ordered or indent != pending_list_indent):
                    flush_list()
                pending_list_ordered = ordered
                pending_list_indent = indent
                pending_list_items.append(text)
                continue

            flush_list()

            if not text:
                continue

            if style_name.startswith("Heading"):
                level_char = style_name[-1]
                level = int(level_char) if level_char.isdigit() else 1
                elements.append(
                    Heading(
                        element_id=stable_id("element", uri, str(len(elements)), "heading"),
                        level=max(1, min(level, 6)),
                        text=text,
                        indent_level=_indent_level(paragraph),
                    )
                )
            else:
                elements.append(
                    TextBlock(
                        element_id=stable_id("element", uri, str(len(elements)), "text"),
                        text=text,
                        indent_level=_indent_level(paragraph),
                    )
                )

        flush_list()
        return elements

    def _convert_tables(
        self, docx_document: DocxDocumentType, uri: str, start_index: int
    ) -> list[DocumentElement]:
        elements: list[DocumentElement] = []
        for offset, table in enumerate(docx_document.tables):
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            elements.append(
                TableElement(
                    element_id=stable_id("element", uri, str(start_index + offset), "table"),
                    rows=rows,
                    extraction_method="rule_based",
                )
            )
        return elements

    async def _convert_images(
        self, docx_document: DocxDocumentType, uri: str, start_index: int
    ) -> list[DocumentElement]:
        picture_parts = []
        for shape in docx_document.inline_shapes:
            if shape.type != WD_INLINE_SHAPE_TYPE.PICTURE:
                continue
            rId = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
            picture_parts.append(docx_document.part.related_parts[rId])

        async def _describe(part) -> str | None:
            try:
                return await self._vision_extractor().describe_figure(part.blob, media_type=part.content_type) or None
            except Exception as exc:  # noqa: BLE001 - degrade this image, not the whole document
                logger.warning("VLM figure description failed on %s: %s", uri, exc)
                return None

        descriptions = await asyncio.gather(*(_describe(part) for part in picture_parts))

        elements: list[DocumentElement] = []
        for description in descriptions:
            elements.append(
                ImageElement(
                    element_id=stable_id("element", uri, str(start_index + len(elements)), "image"),
                    ocr_text=description,
                    extraction_method="vlm" if description is not None else None,
                )
            )
        return elements
