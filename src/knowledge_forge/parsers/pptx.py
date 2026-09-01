import io
import logging

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from knowledge_forge.collectors.base import RawItem
from knowledge_forge.exceptions import ParsingError
from knowledge_forge.ids import stable_id
from knowledge_forge.models.document import Document, DocumentElement, Heading, ImageElement, TableElement, TextBlock
from knowledge_forge.models.source import SourceRef
from knowledge_forge.parsers.base import Parser
from knowledge_forge.vision.extractor import VisionExtractor

logger = logging.getLogger(__name__)


def _format_chart_value(value: float | int | None) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


class PptxParser(Parser):
    """Parses .pptx into a Document using python-pptx.

    Slide index (1-indexed) is used as `page`, consistent with how PDFParser
    numbers pages. Native data tables AND native chart objects are extracted
    deterministically (python-pptx exposes exact cell text and exact
    category/series/value chart data - a VLM could only guess at numbers the
    API already knows for certain, and this environment has no
    slide-rasterizing backend to render a chart as an image anyway). Only
    Picture shapes (embedded images/photos/pre-rendered figures with no
    underlying structured data) go through VisionExtractor.
    """

    supported_extensions = {".pptx"}

    def __init__(self, vision_extractor: VisionExtractor | None = None):
        self._vision = vision_extractor

    def _vision_extractor(self) -> VisionExtractor:
        if self._vision is None:
            self._vision = VisionExtractor()
        return self._vision

    def parse(self, item: RawItem, source_ref: SourceRef) -> Document:
        try:
            elements: list[DocumentElement] = []
            presentation = Presentation(io.BytesIO(item.content))
            for slide_index, slide in enumerate(presentation.slides, start=1):
                elements.extend(self._parse_slide(slide, item.uri, slide_index, len(elements)))
        except Exception as exc:  # noqa: BLE001 - surfaced as a typed parsing error
            raise ParsingError(f"Failed to parse pptx '{item.uri}': {exc}") from exc

        return Document(
            document_id=stable_id("document", item.uri, item.content_hash),
            source=source_ref,
            elements=elements,
        )

    def _parse_chart(self, shape, uri: str, page: int, start_index: int, elements: list[DocumentElement]) -> None:
        chart = shape.chart
        categories = [str(c) for c in chart.plots[0].categories]
        header_row = ["", *categories]
        data_rows = [
            [series.name, *[_format_chart_value(v) for v in series.values]] for series in chart.series
        ]
        rows = [header_row, *data_rows]
        caption = chart.chart_title.text_frame.text if chart.has_title else f"{chart.chart_type.name} chart"
        elements.append(
            TableElement(
                element_id=stable_id("element", uri, str(start_index + len(elements)), "table"),
                page=page,
                rows=rows,
                caption=caption,
                extraction_method="rule_based",
            )
        )

    def _parse_picture(self, shape, uri: str, page: int, start_index: int, elements: list[DocumentElement]) -> None:
        description = None
        try:
            description = self._vision_extractor().describe_figure(shape.image.blob, media_type=shape.image.content_type) or None
        except Exception as exc:  # noqa: BLE001 - degrade this image, not the whole slide/deck
            logger.warning("VLM figure description failed on %s slide %d: %s", uri, page, exc)
        elements.append(
            ImageElement(
                element_id=stable_id("element", uri, str(start_index + len(elements)), "image"),
                page=page,
                ocr_text=description,
                extraction_method="vlm" if description is not None else None,
            )
        )

    def _parse_slide(self, slide, uri: str, page: int, start_index: int) -> list[DocumentElement]:
        elements: list[DocumentElement] = []
        title_shape = slide.shapes.title
        title_text = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        title_shape_id = title_shape.shape_id if title_shape is not None else None

        if title_text:
            elements.append(
                Heading(
                    element_id=stable_id("element", uri, str(start_index + len(elements)), "heading"),
                    page=page,
                    level=1,
                    text=title_text,
                )
            )

        for shape in slide.shapes:
            # Compare by shape_id, not identity - `slide.shapes.title` and the
            # iterator can hand back distinct proxy objects for the same shape.
            if title_shape_id is not None and shape.shape_id == title_shape_id:
                continue

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                self._parse_picture(shape, uri, page, start_index, elements)
                continue

            if shape.has_chart:
                self._parse_chart(shape, uri, page, start_index, elements)
                continue

            if shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                elements.append(
                    TableElement(
                        element_id=stable_id("element", uri, str(start_index + len(elements)), "table"),
                        page=page,
                        rows=rows,
                        extraction_method="rule_based",
                    )
                )
                continue

            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if not text:
                        continue
                    elements.append(
                        TextBlock(
                            element_id=stable_id("element", uri, str(start_index + len(elements)), "text"),
                            page=page,
                            text=text,
                        )
                    )

        return elements
