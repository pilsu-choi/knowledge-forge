"""One-time generator for tests/fixtures/sample.xlsx.

Re-run with `uv run python tests/fixtures/generate_xlsx_fixture.py` if the
fixture needs to be regenerated.
"""

import io
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font
from PIL import Image as PILImage

FIXTURE_PATH = Path(__file__).parent / "sample.xlsx"
REGIONS_FIXTURE_PATH = Path(__file__).parent / "sample_regions.xlsx"


def main() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["id", "name", "amount"])
    sheet.append([1, "Ada", 100])
    sheet.append([2, "Grace", 200])

    img_buf = io.BytesIO()
    PILImage.new("RGB", (80, 50), color=(0, 120, 200)).save(img_buf, format="PNG")
    img_buf.seek(0)
    sheet.add_image(XLImage(img_buf), "E1")

    workbook.save(FIXTURE_PATH)
    print(f"wrote {FIXTURE_PATH}")


def generate_regions_fixture() -> None:
    """A sheet with, top to bottom: a merged title row, a blank separator,
    a 2-level-header table (a merged "2025년" row over "기본급"/"가산율",
    plus a formula cell with no cached value), a blank separator, and a
    merged footnote row - exercises region splitting, merged-cell
    propagation, multi-row header detection and formula extraction all in
    one sheet (see parsers/xlsx.py).
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    bold = Font(bold=True)

    sheet["A1"] = "진료비 기준표"
    sheet.merge_cells("A1:C1")

    sheet["A3"] = "구분"
    sheet["A3"].font = bold
    sheet["B3"] = "2025년"
    sheet["B3"].font = bold
    sheet.merge_cells("B3:C3")

    sheet["B4"] = "기본급"
    sheet["B4"].font = bold
    sheet["C4"] = "가산율"
    sheet["C4"].font = bold

    sheet["A5"] = "외래"
    sheet["B5"] = 10000
    sheet["C5"] = "20%"

    sheet["A6"] = "입원"
    sheet["B6"] = 20000
    sheet["C6"] = "30%"

    sheet["A7"] = "합계"
    # openpyxl never evaluates formulas itself, so a freshly-generated file
    # (unlike one saved by Excel/LibreOffice) has no cached value for this
    # cell - a data_only=True read comes back empty. That's the realistic
    # limit being tested: the formula string is still recoverable even
    # though the computed value isn't.
    sheet["B7"] = "=B5+B6"

    sheet["A9"] = "※ 야간 진료는 별도 가산됨"
    sheet.merge_cells("A9:C9")

    workbook.save(REGIONS_FIXTURE_PATH)
    print(f"wrote {REGIONS_FIXTURE_PATH}")


if __name__ == "__main__":
    main()
    generate_regions_fixture()
