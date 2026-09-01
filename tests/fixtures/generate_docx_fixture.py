"""One-time generator for tests/fixtures/sample.docx.

Re-run with `uv run python tests/fixtures/generate_docx_fixture.py` if the
fixture needs to be regenerated (e.g. after a python-docx upgrade).
"""

from pathlib import Path

from docx import Document

FIXTURE_PATH = Path(__file__).parent / "sample.docx"


def main() -> None:
    doc = Document()
    doc.add_heading("Knowledge Forge Sample", level=1)
    doc.add_paragraph("This is a short paragraph used to exercise the DOCX parser.")
    doc.add_paragraph("first item", style="List Bullet")
    doc.add_paragraph("second item", style="List Bullet")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Role"
    table.cell(1, 0).text = "Ada"
    table.cell(1, 1).text = "Engineer"

    doc.save(FIXTURE_PATH)
    print(f"wrote {FIXTURE_PATH}")


if __name__ == "__main__":
    main()
