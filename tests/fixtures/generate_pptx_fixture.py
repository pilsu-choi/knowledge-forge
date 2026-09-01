"""One-time generator for tests/fixtures/sample.pptx.

Re-run with `uv run python tests/fixtures/generate_pptx_fixture.py` if the
fixture needs to be regenerated.
"""

import io
from pathlib import Path

from PIL import Image as PILImage
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

FIXTURE_PATH = Path(__file__).parent / "sample.pptx"


def main() -> None:
    prs = Presentation()

    title_and_body = prs.slide_layouts[1]
    slide1 = prs.slides.add_slide(title_and_body)
    slide1.shapes.title.text = "Knowledge Forge Sample"
    slide1.placeholders[1].text_frame.text = "This is a short body paragraph used to exercise the PPTX parser."

    blank = prs.slide_layouts[6]
    slide2 = prs.slides.add_slide(blank)
    rows, cols = 2, 2
    table_shape = slide2.shapes.add_table(rows, cols, Inches(1), Inches(1), Inches(4), Inches(1.5))
    table = table_shape.table
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Role"
    table.cell(1, 0).text = "Ada"
    table.cell(1, 1).text = "Engineer"

    slide3 = prs.slides.add_slide(blank)
    img_buf = io.BytesIO()
    PILImage.new("RGB", (80, 50), color=(0, 120, 200)).save(img_buf, format="PNG")
    img_buf.seek(0)
    slide3.shapes.add_picture(img_buf, Inches(1), Inches(1), Inches(2), Inches(1.5))

    slide4 = prs.slides.add_slide(blank)
    chart_data = CategoryChartData()
    chart_data.categories = ["A", "B"]
    chart_data.add_series("Sales", (100, 80))
    slide4.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), chart_data
    )

    prs.save(FIXTURE_PATH)
    print(f"wrote {FIXTURE_PATH}")


if __name__ == "__main__":
    main()
