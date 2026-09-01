import io
from datetime import datetime, timezone
from pathlib import Path

from docx import Document as DocxDocument
from docx.shared import Inches

from knowledge_forge.collectors.base import RawItem
from knowledge_forge.ids import content_hash
from knowledge_forge.models.document import Heading, ImageElement, ListElement, TableElement, TextBlock
from knowledge_forge.models.source import SourceRef, SourceType
from knowledge_forge.parsers.docx import DocxParser
from knowledge_forge.vision.extractor import VisionExtractor

FIXTURE = Path(__file__).parent.parent / "fixtures" / "sample.docx"

_FIGURE_DESCRIPTION = "A solid blue rectangle placeholder image."


class _FakeVisionExtractor(VisionExtractor):
    def __init__(self):
        self.figure_calls = 0

    async def describe_figure(self, image_bytes: bytes, media_type: str = "image/png") -> str:
        self.figure_calls += 1
        return _FIGURE_DESCRIPTION


async def test_docx_parses_heading_paragraph_list_table_and_picture():
    data = FIXTURE.read_bytes()
    item = RawItem(
        uri=f"file://{FIXTURE}",
        content=data,
        content_hash=content_hash(data),
        modified_at=datetime.now(timezone.utc),
    )
    source_ref = SourceRef(source_id="s1", source_type=SourceType.LOCAL_FILE, uri=item.uri)

    vision = _FakeVisionExtractor()
    document = await DocxParser(vision_extractor=vision).parse(item, source_ref)

    headings = [e for e in document.elements if isinstance(e, Heading)]
    texts = [e for e in document.elements if isinstance(e, TextBlock)]
    lists = [e for e in document.elements if isinstance(e, ListElement)]
    tables = [e for e in document.elements if isinstance(e, TableElement)]
    images = [e for e in document.elements if isinstance(e, ImageElement)]

    assert len(headings) == 1
    assert headings[0].level == 1
    assert headings[0].text == "Knowledge Forge Sample"

    assert len(texts) == 1
    assert "short paragraph" in texts[0].text

    assert len(lists) == 1
    assert lists[0].items == ["first item", "second item"]

    assert len(tables) == 1
    assert tables[0].rows[0] == ["Name", "Role"]
    assert tables[0].rows[1] == ["Ada", "Engineer"]
    assert tables[0].extraction_method == "rule_based"

    # the embedded picture is described via VLM, not left unparsed.
    assert len(images) == 1
    assert images[0].extraction_method == "vlm"
    assert images[0].ocr_text == _FIGURE_DESCRIPTION
    assert vision.figure_calls == 1

    assert all(e.page is None for e in document.elements)


async def test_docx_captures_paragraph_and_nested_list_indent_levels():
    docx_document = DocxDocument()
    docx_document.add_paragraph("top level")
    indented = docx_document.add_paragraph("indented once")
    indented.paragraph_format.left_indent = Inches(0.5)
    docx_document.add_paragraph("top bullet", style="List Bullet")
    nested_bullet = docx_document.add_paragraph("nested bullet", style="List Bullet 2")
    nested_bullet.paragraph_format.left_indent = Inches(1.0)

    buf = io.BytesIO()
    docx_document.save(buf)
    data = buf.getvalue()
    item = RawItem(
        uri="file://in-memory.docx",
        content=data,
        content_hash=content_hash(data),
        modified_at=datetime.now(timezone.utc),
    )
    source_ref = SourceRef(source_id="s1", source_type=SourceType.LOCAL_FILE, uri=item.uri)

    document = await DocxParser(vision_extractor=_FakeVisionExtractor()).parse(item, source_ref)

    texts = [e for e in document.elements if isinstance(e, TextBlock)]
    lists = [e for e in document.elements if isinstance(e, ListElement)]

    assert texts[0].text == "top level"
    assert texts[0].indent_level is None
    assert texts[1].text == "indented once"
    assert texts[1].indent_level == 1

    # A change in nesting depth starts a new ListElement rather than
    # collapsing the sub-item into the outer list's text.
    assert len(lists) == 2
    assert lists[0].items == ["top bullet"]
    assert lists[0].indent_level is None
    assert lists[1].items == ["nested bullet"]
    assert lists[1].indent_level == 2
